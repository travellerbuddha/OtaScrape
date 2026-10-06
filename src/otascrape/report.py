from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .compare import Comparison, PriceChange
from .config import Config
from .models import Offer

BOARD_LABELS = {
    "RO": "Sadece oda", "BB": "Oda kahvaltı", "HB": "Yarım pansiyon", "FB": "Tam pansiyon",
    "AI": "Her şey dahil", "UAI": "Ultra her şey dahil", "UNKNOWN": "Belirsiz",
}
_RED = PatternFill("solid", fgColor="F8D7DA")
_GREEN = PatternFill("solid", fgColor="D4EDDA")
_HEAD = PatternFill("solid", fgColor="1F3A5F")


POLICY_LABELS = {"free": "ücretsiz iptal", "nofree": "ücretsiz iptal yok", "unknown": "iptal koşulu bilinmiyor"}


def _group_title(c: Comparison) -> str:
    extra = [x for x in (c.room_type, POLICY_LABELS.get(c.cancel_policy)) if x]
    return " · ".join([BOARD_LABELS.get(c.board, c.board), *extra])


def _hotel_names(cfg: Config) -> dict[str, str]:
    return {h.id: h.name for h in cfg.hotels}


def _cancel(v: bool | None) -> str:
    return {True: "Ücretsiz iptal", False: "Ücretsiz iptal yok"}.get(v, "?")


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, 1):
        width = widths[i - 1] if widths else max(12, min(40, len(header) + 4))
        ws.column_dimensions[get_column_letter(i)].width = width


def write_excel(path: Path, cfg: Config, offers: list[Offer], comparisons: list[Comparison],
                changes: list[PriceChange], errors: list) -> None:
    names = _hotel_names(cfg)
    cur = cfg.base_currency
    wb = Workbook()

    ws = wb.active
    ws.title = "Karşılaştırma"
    heads = ["Otel", "Giriş", "Konaklama", "Pansiyon", "Oda tipi (eşleşen)", "Acenta / Satıcı", "Kanal", "Oda", "İptal",
             f"Toplam ({cur})", f"Gecelik ({cur})", f"Kişi başı gecelik ({cur})",
             f"Referansa fark ({cur})", "Referansa fark %", "Referans", "En ucuz", "Parite ihlali"]
    rows, marks = [], []
    for c in comparisons:
        for r in c.rows:
            rows.append([
                names.get(c.hotel_id, c.hotel_id), c.check_in, c.stay_name, BOARD_LABELS.get(c.board, c.board), c.room_type or "-",
                r.seller, r.channel, r.room_name, _cancel(r.free_cancellation),
                round(r.total, 2), round(r.per_night, 2), round(r.per_person_night, 2),
                round(r.diff_abs, 2), round(r.diff_pct, 2) / 100,
                "✔" if r.is_reference else "", "✔" if r.is_cheapest else "", "⚠" if r.parity_breach else "",
            ])
            marks.append((r.parity_breach, r.is_cheapest))
    _sheet(ws, heads, rows, [24, 12, 12, 16, 18, 22, 12, 34, 18, 14, 14, 18, 16, 14, 10, 10, 12])
    for idx, (breach, cheapest) in enumerate(marks, 2):
        fill = _RED if breach else (_GREEN if cheapest else None)
        for cell in ws[idx]:
            if fill:
                cell.fill = fill
        ws.cell(idx, 2).number_format = "yyyy-mm-dd"
        for col in (10, 11, 12, 13):
            ws.cell(idx, col).number_format = "#,##0.00"
        ws.cell(idx, 14).number_format = "0.00%"

    ws = wb.create_sheet("Parite İhlalleri")
    breach_rows = []
    for c in comparisons:
        for r in c.rows:
            if r.parity_breach:
                breach_rows.append([
                    names.get(c.hotel_id, c.hotel_id), c.check_in, c.stay_name, _group_title(c),
                    r.seller, round(c.reference_total, 2), round(r.total, 2), round(r.diff_abs, 2), round(r.diff_pct, 2) / 100,
                ])
    breach_rows.sort(key=lambda x: x[8])
    _sheet(ws, ["Otel", "Giriş", "Konaklama", "Grup (pansiyon · oda · iptal)", "Ucuz satan acenta", f"Direkt fiyat ({cur})",
                f"Acenta fiyatı ({cur})", f"Fark ({cur})", "Fark %"], breach_rows, [24, 12, 12, 16, 24, 16, 16, 14, 10])
    for idx in range(2, len(breach_rows) + 2):
        ws.cell(idx, 2).number_format = "yyyy-mm-dd"
        ws.cell(idx, 9).number_format = "0.00%"

    ws = wb.create_sheet("Fiyat Değişimi")
    _sheet(ws, ["Otel", "Giriş", "Konaklama", "Pansiyon", "Acenta", f"Önceki ({cur})", f"Şimdi ({cur})", f"Fark ({cur})", "Fark %"],
           [[names.get(x.hotel_id, x.hotel_id), x.check_in, x.stay_name, BOARD_LABELS.get(x.board, x.board), x.seller,
             round(x.previous_total, 2), round(x.current_total, 2), round(x.diff_abs, 2), round(x.diff_pct, 2) / 100] for x in changes],
           [24, 12, 12, 16, 24, 14, 14, 14, 10])
    for idx in range(2, len(changes) + 2):
        ws.cell(idx, 2).number_format = "yyyy-mm-dd"
        ws.cell(idx, 9).number_format = "0.00%"

    ws = wb.create_sheet("Ham Teklifler")
    _sheet(ws, ["Otel", "Kanal", "Acenta / Satıcı", "Konaklama", "Giriş", "Gece", "Yetişkin", "Çocuk", "Oda", "Pansiyon",
                "Toplam", "Para birimi", "İptal", "Vergi dahil", "Çekim zamanı", "URL"],
           [[names.get(o.hotel_id, o.hotel_id), o.channel, o.seller, o.stay_name, o.check_in, o.nights, o.adults, o.children,
             o.room_name, BOARD_LABELS.get(o.board, o.board), o.total_price, o.currency, _cancel(o.free_cancellation),
             {True: "Evet", False: "Hayır"}.get(o.taxes_included, "?"), o.scraped_at.strftime("%Y-%m-%d %H:%M"), o.source_url] for o in offers],
           [24, 12, 22, 12, 12, 7, 9, 7, 34, 16, 12, 10, 14, 11, 17, 50])

    ws = wb.create_sheet("Hatalar")
    _sheet(ws, ["Otel", "Kanal", "Giriş", "Konaklama", "Mesaj", "Zaman"],
           [[names.get(e["hotel_id"], e["hotel_id"]), e["channel"], e["check_in"], e["stay_name"], e["message"], e["occurred_at"]] for e in errors],
           [24, 12, 12, 12, 80, 22])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


_CSS = """
body{font-family:system-ui,Segoe UI,Arial,sans-serif;margin:24px;color:#1c2733;background:#f6f8fa}
h1{margin:0 0 4px}h2{margin:28px 0 8px;border-bottom:2px solid #1f3a5f;padding-bottom:4px}
h3{margin:18px 0 6px;font-size:15px}.meta{color:#5b6b7b;font-size:13px}
table{border-collapse:collapse;background:#fff;margin:6px 0 14px;font-size:13px}
th{background:#1f3a5f;color:#fff;padding:6px 10px;text-align:left}td{padding:5px 10px;border-bottom:1px solid #e3e8ee}
td.n{text-align:right;font-variant-numeric:tabular-nums}tr.breach td{background:#f8d7da}tr.cheapest td{background:#e3f4e8}
.kpi{display:inline-block;background:#fff;border:1px solid #d5dde5;border-radius:8px;padding:10px 16px;margin:0 10px 10px 0}
.kpi b{display:block;font-size:22px}.err{color:#a4262c}
"""


def _fmt(v: float) -> str:
    return f"{v:,.2f}"


def write_html(path: Path, cfg: Config, comparisons: list[Comparison], changes: list[PriceChange],
               errors: list, generated_at: datetime) -> None:
    names = _hotel_names(cfg)
    cur = cfg.base_currency
    breaches = [(c, r) for c in comparisons for r in c.rows if r.parity_breach]
    out = [f"<!doctype html><html lang='tr'><meta charset='utf-8'><title>OTA Fiyat Karşılaştırma</title><style>{_CSS}</style><body>",
           "<h1>OTA Fiyat Karşılaştırma</h1>",
           f"<div class='meta'>Oluşturma: {generated_at:%Y-%m-%d %H:%M} · Para birimi: {escape(cur)} · "
           f"Kişi başı hesap: {'yalnız yetişkin' if cfg.pax_basis == 'adults' else 'tüm misafirler'} · "
           f"Parite toleransı: %{cfg.parity_tolerance_pct:g}</div>",
           f"<p><span class='kpi'><b>{len(comparisons)}</b>karşılaştırma</span>"
           f"<span class='kpi'><b>{len(breaches)}</b>parite ihlali</span>"
           f"<span class='kpi'><b>{len(changes)}</b>fiyat değişimi</span>"
           f"<span class='kpi'><b class='{'err' if errors else ''}'>{len(errors)}</b>hata</span></p>"]

    if not any(h.room_types for h in cfg.hotels):
        out.append("<p class='meta' style='color:#a4262c'>⚠ Oda tipi eşlemesi (room_types) tanımlı değil: her satıcının <b>en ucuz odası</b> "
                   "kıyaslanıyor, farklı oda tipleri karışabilir. Eşleme için config.example.yaml'a bakın.</p>")
    out.append("<h2>Parite ihlalleri (acenta direkt fiyatın altında)</h2>")
    if breaches:
        out.append(f"<table><tr><th>Otel</th><th>Giriş</th><th>Konaklama</th><th>Grup</th><th>Acenta</th><th>Direkt ({cur})</th><th>Acenta ({cur})</th><th>Fark %</th></tr>")
        for c, r in sorted(breaches, key=lambda x: x[1].diff_pct):
            out.append(f"<tr class='breach'><td>{escape(names.get(c.hotel_id, c.hotel_id))}</td><td>{c.check_in}</td><td>{escape(c.stay_name)}</td>"
                       f"<td>{escape(_group_title(c))}</td><td>{escape(r.seller)}</td><td class='n'>{_fmt(c.reference_total)}</td>"
                       f"<td class='n'>{_fmt(r.total)}</td><td class='n'>{r.diff_pct:+.1f}%</td></tr>")
        out.append("</table>")
    else:
        out.append("<p>İhlal yok.</p>")

    out.append("<h2>Karşılaştırmalar</h2>")
    for c in comparisons:
        out.append(f"<h3>{escape(names.get(c.hotel_id, c.hotel_id))} · {c.check_in} · {escape(c.stay_name)} · "
                   f"{escape(_group_title(c))} <span class='meta'>(referans: {escape(c.reference_seller)}, "
                   f"yayılım %{c.spread_pct:.1f})</span></h3>")
        out.append(f"<table><tr><th>Acenta</th><th>Kanal</th><th>İptal</th><th>Toplam ({cur})</th><th>Gecelik</th><th>Kişi başı gecelik</th><th>Referansa fark</th><th>%</th></tr>")
        for r in c.rows:
            cls = "breach" if r.parity_breach else ("cheapest" if r.is_cheapest else "")
            out.append(f"<tr class='{cls}'><td>{escape(r.seller)}{' ★' if r.is_reference else ''}</td><td>{escape(r.channel)}</td>"
                       f"<td>{_cancel(r.free_cancellation)}</td><td class='n'>{_fmt(r.total)}</td><td class='n'>{_fmt(r.per_night)}</td>"
                       f"<td class='n'>{_fmt(r.per_person_night)}</td><td class='n'>{r.diff_abs:+,.2f}</td><td class='n'>{r.diff_pct:+.1f}%</td></tr>")
        out.append("</table>")

    if changes:
        out.append("<h2>Önceki taramaya göre fiyat değişimi</h2>")
        out.append(f"<table><tr><th>Otel</th><th>Giriş</th><th>Konaklama</th><th>Pansiyon</th><th>Acenta</th><th>Önceki</th><th>Şimdi</th><th>%</th></tr>")
        for x in changes[:100]:
            out.append(f"<tr><td>{escape(names.get(x.hotel_id, x.hotel_id))}</td><td>{x.check_in}</td><td>{escape(x.stay_name)}</td>"
                       f"<td>{BOARD_LABELS.get(x.board, x.board)}</td><td>{escape(x.seller)}</td><td class='n'>{_fmt(x.previous_total)}</td>"
                       f"<td class='n'>{_fmt(x.current_total)}</td><td class='n'>{x.diff_pct:+.1f}%</td></tr>")
        out.append("</table>")

    if errors:
        out.append("<h2>Hatalar</h2><table><tr><th>Otel</th><th>Kanal</th><th>Giriş</th><th>Konaklama</th><th>Mesaj</th></tr>")
        for e in errors:
            out.append(f"<tr><td>{escape(names.get(e['hotel_id'], e['hotel_id']))}</td><td>{escape(e['channel'])}</td>"
                       f"<td>{e['check_in']}</td><td>{escape(e['stay_name'])}</td><td class='err'>{escape(e['message'])}</td></tr>")
        out.append("</table>")
    out.append("</body></html>")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(out), encoding="utf-8")
