from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from otascrape import db
from otascrape.adapters.base import BlockedError, ScrapeError
from otascrape.adapters.booking import build_url, parse_offers
from otascrape.cli import main
from otascrape.models import Search, Stay

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture
def search():
    return Search("h1", "Hotel 1", "booking", "https://www.booking.com/hotel/tr/x.html?foo=1", date(2026, 11, 1),
                  Stay("7N-2AD-1CH", 7, 2, 1, (7,)), "EUR")


def test_build_url(search):
    url = build_url(search)
    for part in ("checkin=2026-11-01", "checkout=2026-11-08", "group_adults=2", "group_children=1", "age=7", "selected_currency=EUR"):
        assert part in url
    assert "foo=1" not in url and url.startswith("https://www.booking.com/hotel/tr/x.html?")


def test_parse_offers(search):
    offers = parse_offers((FIX / "booking_rows.html").read_text(), search)
    assert [(o.room_name, o.total_price, o.currency, o.board, o.free_cancellation, o.taxes_included) for o in offers] == [
        ("Deluxe Room", 2345, "EUR", "AI", True, False),
        ("Deluxe Room", 2010, "EUR", "BB", False, None),
        ("Family Suite", 3980.5, "EUR", "UAI", None, True),
    ]
    assert all(o.seller == "Booking.com" and o.nights == 7 for o in offers)


def test_parse_soldout_and_blocked(search):
    assert parse_offers((FIX / "booking_soldout.html").read_text(), search) == []
    with pytest.raises(BlockedError):
        parse_offers((FIX / "booking_blocked.html").read_text(), search)
    with pytest.raises(ScrapeError):
        parse_offers("<html><body>garip bir sayfa</body></html>", search)


def test_mock_end_to_end(tmp_path, monkeypatch):
    import yaml
    cfg_path = tmp_path / "c.yaml"
    cfg_path.write_text(yaml.safe_dump({
        "database": str(tmp_path / "t.db"), "output_dir": str(tmp_path / "out"),
        "direct_sellers": ["official site"], "fx": {"EUR": 1},
        "stays": [{"name": "7N-2AD", "nights": 7, "adults": 2}, {"name": "3N-2AD", "nights": 3, "adults": 2}],
        "checkin": {"offsets_days": [14, 30]},
        "hotels": [
            {"id": "a", "name": "A", "channels": {"trivago": "u", "booking": "u"}},
            {"id": "b", "name": "B", "channels": {"check24": "u", "expedia": "u"}},
        ],
    }))
    assert main(["-c", str(cfg_path), "run", "--mock"]) == 0
    assert main(["-c", str(cfg_path), "run", "--mock"]) == 0  # ikinci tarama -> değişim hesaplanabilir (aynı sahte veri: değişim yok)
    xlsx = sorted((tmp_path / "out").glob("*.xlsx"))[-1]
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Karşılaştırma", "Parite İhlalleri", "Fiyat Değişimi", "Ham Teklifler", "Hatalar"]
    assert wb["Karşılaştırma"].max_row > 20 and wb["Ham Teklifler"].max_row > 20
    assert list((tmp_path / "out").glob("*.html"))
    conn = db.connect(tmp_path / "t.db")
    assert db.latest_run_id(conn) == 2 and db.latest_run_id(conn, before=2) == 1


def test_unknown_channel_is_logged_not_fatal(tmp_path):
    import yaml
    p = tmp_path / "c.yaml"
    p.write_text(yaml.safe_dump({
        "database": str(tmp_path / "t.db"), "output_dir": str(tmp_path / "out"),
        "stays": [{"nights": 7, "adults": 2}], "checkin": {"offsets_days": [14]},
        "hotels": [{"id": "a", "channels": {"kayak": "u"}}],
    }))
    assert main(["-c", str(p), "run"]) == 1  # hata var -> çıkış kodu 1
    conn = db.connect(tmp_path / "t.db")
    err = db.load_errors(conn, 1)
    assert len(err) == 1 and "Bilinmeyen kanal" in err[0]["message"]


def test_booking_real_page_with_waf_script_is_not_blocked(search):
    # Gerçek yakalamada görüldü: sayfa 1,8 MB, JS içinde 'awswaf' geçiyor, tablo yok.
    big = "<html><body><p>Mardan Palace</p>" + "<div>x</div>" * 5000 + "<script>var w='awswaf'; var m='sold out';</script></body></html>"
    with pytest.raises(ScrapeError) as exc:
        parse_offers(big, search)
    assert not isinstance(exc.value, BlockedError)


def test_booking_dates_not_applied(search):
    html = "<html><body><div>Select dates to see this property's availability and prices</div></body></html>"
    with pytest.raises(ScrapeError, match="tarih"):
        parse_offers(html, search)
